"""GAP-089: the one xlsx reader both spreadsheet importers share.

Cells arrive typed from openpyxl; the importers want strings (plus real dates),
so every scalar is coerced here in one place:

- ``int`` / integral ``float`` → ``"3"`` (Excel stores number-typed phone cells
  as floats — ``447985414214.0`` must not become ``"447985414214.0"``);
- other floats → ``str()``; ``bool`` → ``"True"``/``"False"``;
- ``str`` → stripped, ``""`` → ``None``;
- ``datetime`` / ``date`` → left alone (``parse_sheet_date`` normalises later);
- ``None`` → ``None``.

Rows whose every cell coerces to ``None`` are skipped (hand-kept sheets have
blank spacer rows); reading stops after ``BLANK_RUN_LIMIT`` of them in a row,
because the real exports report ~1M rows — trailing formatting keeps the used
range open.
Unnamed header cells get a positional ``col_<n>`` key (1-based) so a stray
mail-merge column is preserved rather than colliding on ``None``; a repeated
header name gets a ``<name>_<n>`` suffix for the same reason.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

#: Consecutive all-blank rows that end the sheet (see module doc).
BLANK_RUN_LIMIT = 200


def _coerce(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return f"{value:.0f}" if value.is_integer() else str(value)
    if isinstance(value, datetime | date):
        return value
    text = str(value).strip()
    return text or None


def read_sheet(path: str | Path, sheet_name: str) -> list[dict[str, Any]]:
    """Return the named sheet as a list of header-keyed dicts (see module doc).

    Raises ``KeyError`` for an unknown sheet name (openpyxl's own error).
    """
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_header = next(rows_iter)
        except StopIteration:
            return []
        header: list[str] = []
        for i, cell in enumerate(raw_header, start=1):
            name = str(cell).strip() if cell is not None and str(cell).strip() else f"col_{i}"
            # A repeated header would otherwise collapse to the last column's
            # value in the row dict; suffix the repeat with its position.
            header.append(name if name not in header else f"{name}_{i}")
        rows: list[dict[str, Any]] = []
        blank_run = 0
        for raw in rows_iter:
            values = [_coerce(cell) for cell in raw]
            if all(v is None for v in values):
                # A blank spacer row is skipped; only a long run of them means
                # the used range has run out (the real exports report ~1M rows).
                blank_run += 1
                if blank_run >= BLANK_RUN_LIMIT:
                    break
                continue
            blank_run = 0
            # Trailing cells beyond the header (or short rows) are tolerated.
            values = (values + [None] * len(header))[: len(header)]
            rows.append(dict(zip(header, values, strict=True)))
        return rows
    finally:
        workbook.close()
