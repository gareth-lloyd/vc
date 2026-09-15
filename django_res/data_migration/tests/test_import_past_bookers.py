"""GAP-089: `import_past_bookers` — Contacts → Person, Booking History → PastStay."""

from __future__ import annotations

import re
from io import StringIO
from pathlib import Path
from typing import Any

import pytest
from django.core.management import CommandError, call_command
from openpyxl import Workbook

from accounts.models import Person, PersonEmail
from properties.models import Country, Property, Region
from reservations.models import PastStay

pytestmark = pytest.mark.django_db

CONTACT_HEADER = [
    None,
    "Notes notes",
    "Agent / Advisor",
    "Prefix",
    "First Name",
    "Last Name",
    "Email",
    "Mailing Street",
    "Mailing City",
    "Mailing Zip",
    "Mailing Country",
    "No. of Bookings",
    "Lead Source",
]
HISTORY_HEADER = [
    "First Name",
    "Last Name",
    "Booking Number",
    "Villa Booked",
    "Destination",
    "Year of Check-In",
]


def _workbook(
    path: Path, contacts: list[list[Any]], history: list[list[Any]], *, trailing_blank: int = 3
) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Contacts"
    ws.append(CONTACT_HEADER)
    for row in contacts:
        ws.append(row)
    for _ in range(trailing_blank):
        ws.append([None] * len(CONTACT_HEADER))
    hs = wb.create_sheet("Booking History")
    hs.append(HISTORY_HEADER)
    for row in history:
        hs.append(row)
    wb.save(path)
    return path


def _contact(
    first: str,
    last: str,
    email: str | None = "ada@example.com",
    *,
    prefix: str = "Ms",
    notes: str | None = None,
    advisor: str | None = None,
    country: str = "UK",
) -> list[Any]:
    return [
        None,
        notes,
        advisor,
        prefix,
        first,
        last,
        email,
        "1 Test St",
        "Bath",
        "BA1 1AA",
        country,
        1,
        "Villa Collective",
    ]


def _run(path: Path, *args: str) -> str:
    out = StringIO()
    call_command("import_past_bookers", "--file", str(path), *args, stdout=out)
    return out.getvalue()


@pytest.fixture
def villa(db: None) -> Property:
    country, _ = Country.objects.get_or_create(
        iso2="GR", defaults={"name": "Greece", "iso3": "GRC"}
    )
    region = Region.objects.create(country=country, name="Corfu", slug="corfu")
    return Property.objects.create(
        name="Villa Yeraki", display_name="Villa Yeraki", slug="villa-yeraki", region=region
    )


def test_creates_person_with_address_and_linked_past_stay(tmp_path: Path, villa: Property) -> None:
    Country.objects.get_or_create(iso2="GB", defaults={"name": "United Kingdom", "iso3": "GBR"})
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace", notes="Loves Corfu", advisor="Bob (10%)")],
        [["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019]],
    )

    out = _run(path)

    person = Person.objects.get(last_name="Lovelace")
    assert person.legacy_id is not None and person.legacy_id.startswith("sheet-person-")
    assert person.title == "Ms"
    assert person.address_line_1 == "1 Test St"
    assert person.town == "Bath"
    assert person.post_code == "BA1 1AA"
    assert person.country is not None and person.country.iso2 == "GB"
    assert person.kind == "customer"
    assert person.notes == "Loves Corfu\nAgent/advisor: Bob (10%)"
    assert list(person.emails.values_list("email", flat=True)) == ["ada@example.com"]
    stay = PastStay.objects.get()
    assert stay.person == person
    assert stay.booking_number == "BN123"
    assert stay.villa_name == "Yeraki"
    assert stay.property == villa
    assert stay.destination == "Corfu"
    assert stay.year == 2019
    assert stay.legacy_id is not None and stay.legacy_id.startswith("sheet-stay-")
    assert "person" in out and "past_stay" in out


def test_spouses_sharing_an_email_become_two_people(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [
            _contact("Orlando", "Fraser", "fraser@example.com"),
            _contact("Clemmie", "Fraser-Jones", "fraser@example.com"),
        ],
        [],
    )

    _run(path)

    assert Person.objects.filter(emails__email="fraser@example.com").count() == 2
    assert PersonEmail.objects.filter(email="fraser@example.com").count() == 2


def test_rerun_creates_nothing_and_keeps_operator_edits(tmp_path: Path, villa: Property) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace", notes="Loves Corfu")],
        [["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019]],
    )
    _run(path)
    person = Person.objects.get(last_name="Lovelace")
    person.town = "Bristol"
    person.save(update_fields=["town"])

    out = _run(path)

    person.refresh_from_db()
    assert person.town == "Bristol"
    assert person.notes == "Loves Corfu"
    assert Person.objects.count() == 1
    assert PersonEmail.objects.count() == 1
    assert PastStay.objects.count() == 1
    assert "exists" in out
    assert PastStay.objects.get().person == person


def test_history_links_to_existing_customer_when_not_in_contacts(tmp_path: Path) -> None:
    existing = Person.objects.create(first_name="Grace", last_name="Hopper", kind="customer")
    path = _workbook(tmp_path / "b.xlsx", [], [["Grace", "hopper", "BN9", "Nowhere", "", 2018]])

    out = _run(path)

    stay = PastStay.objects.get()
    assert stay.person == existing
    assert stay.property is None
    assert "Nowhere" in out  # unmatched villa is reported, not guessed


def test_unmatched_person_is_skipped_and_reported(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "b.xlsx", [], [["No", "Body", "BN1", "Yeraki", "Corfu", 2019]])

    out = _run(path)

    assert PastStay.objects.count() == 0
    assert "person_unmatched" in out
    assert "No Body" in out


def test_ambiguous_sheet_name_is_skipped_in_history(tmp_path: Path) -> None:
    # Two contacts with the same name → a stay by that name cannot be attributed.
    path = _workbook(
        tmp_path / "b.xlsx",
        [
            _contact("Sam", "Smith", "sam1@example.com"),
            _contact("Sam", "Smith", "sam2@example.com"),
        ],
        [["Sam", "Smith", "BN1", "Yeraki", "Corfu", 2019]],
    )

    out = _run(path)

    assert Person.objects.filter(last_name="Smith").count() == 2
    assert PastStay.objects.count() == 0
    assert "person_unmatched" in out


def test_multiline_booking_number_and_bad_year(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace")],
        [["Ada", "Lovelace", "BN40\nCANCELLED", "Yeraki", "Corfu", 1905]],
    )

    out = _run(path)

    stay = PastStay.objects.get()
    assert stay.booking_number == "BN40"
    assert stay.year is None
    assert stay.notes == "CANCELLED\nYear as written: 1905"
    assert "bad_year" in out


def test_dry_run_writes_nothing_and_prints_report(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace")],
        [["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019]],
    )

    out = _run(path, "--dry-run")

    assert "[dry-run]" in out
    assert "past_stay" in out
    assert Person.objects.count() == 0
    assert PastStay.objects.count() == 0


def test_missing_file_is_a_command_error(tmp_path: Path) -> None:
    with pytest.raises(CommandError, match="No such file"):
        _run(tmp_path / "nope.xlsx")


def test_missing_sheet_is_a_command_error(tmp_path: Path) -> None:
    wb = Workbook()
    path = tmp_path / "wrong.xlsx"
    wb.save(path)

    with pytest.raises(CommandError, match="missing a sheet"):
        _run(path)


def test_one_bad_row_does_not_abort_the_run(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [
            _contact("", "", None),  # blank → skipped
            _contact("Bad", "Email", "not-an-email"),
            _contact("Ada", "Lovelace"),
        ],
        [],
    )

    out = _run(path)

    assert Person.objects.filter(last_name__in=["Email", "Lovelace"]).count() == 2
    assert Person.objects.get(last_name="Email").emails.count() == 0
    assert "blank_row" in out
    assert "invalid_email" in out


def test_identical_history_rows_in_one_run_are_reported_as_duplicates(tmp_path: Path) -> None:
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace")],
        [
            ["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019],
            ["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019],
            ["Ada", "Lovelace", "BN124", "Yeraki", "Corfu", 2019],
        ],
    )

    out = _run(path)

    assert PastStay.objects.count() == 2
    assert "duplicate_row" in out
    assert "exists" not in out


def test_a_contact_that_fails_after_its_person_is_counted_counts_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """BUG-030 §34: same over-count shape as the enquiry sheet."""
    from data_migration.management.commands import import_past_bookers

    def _boom(*args: Any, **kwargs: Any) -> bool:
        raise RuntimeError("note write failed")

    monkeypatch.setattr(import_past_bookers, "append_note_line", _boom)
    path = _workbook(tmp_path / "b.xlsx", [_contact("Ada", "Lovelace", notes="VIP")], [])

    out = _run(path)

    assert Person.objects.count() == 0
    assert "note write failed" in out
    assert not re.search(r"created\s+person", out)


def test_a_stay_whose_write_failed_is_not_treated_as_a_duplicate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """BUG-030 §34: a rolled-back stay must not stay in the run's `seen` set,
    or an identical later row is skipped as a duplicate and never written."""
    real_get_or_create = PastStay.objects.get_or_create
    calls = {"n": 0}

    def _flaky(**kwargs: Any) -> Any:
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("stay write failed")
        return real_get_or_create(**kwargs)

    monkeypatch.setattr(PastStay.objects, "get_or_create", _flaky)
    path = _workbook(
        tmp_path / "b.xlsx",
        [_contact("Ada", "Lovelace")],
        [
            ["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019],
            ["Ada", "Lovelace", "BN123", "Yeraki", "Corfu", 2019],
        ],
    )

    out = _run(path)

    assert PastStay.objects.count() == 1
    assert "stay write failed" in out
    assert "duplicate_row" not in out
