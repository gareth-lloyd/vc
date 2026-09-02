"""GAP-089: `read_sheet` — the one xlsx reader both spreadsheet importers share."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from data_migration.sheets.xlsx import read_sheet


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Contacts"
    # Col A is an unnamed mail-merge column in the real export.
    ws.append([None, "First Name", "Phone", "Enquiry Date", "Count", "Flag"])
    ws.append(["sent", "  Ada ", 447985414214.0, datetime(2019, 6, 1, 12, 30), 3, True])
    ws.append([None, "Bob", "+44 7000", "2020-01-02", 2.5, None])
    ws.append([None, "", None, None, None, None])  # only-empty-strings row → blank
    ws.append([None, None, None, None, None, None])
    ws.append([None, "Ghost", None, None, None, None])  # after the blank row: ignored
    wb.create_sheet("Other").append(["x"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def test_read_sheet_coerces_cells_and_stops_at_first_blank_row(workbook: Path) -> None:
    rows = read_sheet(workbook, "Contacts")

    assert [r["First Name"] for r in rows] == ["Ada", "Bob"]
    # Integral floats (Excel's number-typed phone cells) drop the ".0".
    assert rows[0]["Phone"] == "447985414214"
    assert rows[1]["Phone"] == "+44 7000"
    # Dates stay dates; ISO strings stay strings (the caller parses).
    assert rows[0]["Enquiry Date"] == datetime(2019, 6, 1, 12, 30)
    assert rows[1]["Enquiry Date"] == "2020-01-02"
    assert rows[0]["Count"] == "3"
    assert rows[1]["Count"] == "2.5"
    assert rows[0]["Flag"] == "True"
    assert rows[1]["Flag"] is None
    # The unnamed column gets a positional key rather than being dropped.
    assert rows[0]["col_1"] == "sent"
    assert rows[1]["col_1"] is None


def test_read_sheet_unknown_sheet_raises(workbook: Path) -> None:
    with pytest.raises(KeyError):
        read_sheet(workbook, "Nope")


def test_read_sheet_keeps_date_cells(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["When"])
    ws.append([date(2018, 3, 4)])
    path = tmp_path / "d.xlsx"
    wb.save(path)

    (row,) = read_sheet(path, ws.title)

    assert isinstance(row["When"], datetime | date)
    assert (row["When"].date() if isinstance(row["When"], datetime) else row["When"]) == date(
        2018, 3, 4
    )
